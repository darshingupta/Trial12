"""Build the deliverables from paper.py: the online test page, the Word paper and the Word answer key."""
import json, base64, sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from paper import P, CONVENTIONS

OUT = sys.argv[1] if len(sys.argv) > 1 else 'out'
os.makedirs(OUT, exist_ok=True)
LET = 'ABCD'
ROMAN = ['I', 'II', 'III', 'IV', 'V', 'VI']
INK = RGBColor(0x18, 0x21, 0x2C); ACC = RGBColor(0x2C, 0x5A, 0x86); MUT = RGBColor(0x5A, 0x65, 0x73)

# ---------------------------------------------------------------- HTML
def build_html():
    data = {'ver': 'r%d' % sum(len(q['stem']) for p in P for q in p['qs']), 'conventions': CONVENTIONS,
            'passages': [{'title': p['title'], 'kind': p['kind'], 'intro': p['intro'], 'conds': p['conds'],
                          'qs': [{'stem': q['stem'], 'opts': q['opts']} for q in p['qs']]} for p in P]}
    key = {'q': [{'a': q['ans'], 'why': q['why'], 'trap': q['trap']} for p in P for q in p['qs']],
           'sol': [p['solution'] for p in P]}
    enc = base64.b64encode(json.dumps(key).encode('utf-8')).decode()[::-1]
    t = open('template.html').read()
    t = t.replace('__DATA__', json.dumps(data, ensure_ascii=False)).replace('__KEY__', json.dumps(enc))
    open(os.path.join(OUT, 'ar_gauntlet.html'), 'w').write(t)

# ---------------------------------------------------------------- Word helpers
def base_doc():
    d = Document()
    s = d.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles['Normal']; st.font.name = 'Cambria'; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Cambria')
    st.paragraph_format.space_after = Pt(4)
    for name, size in (('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)):
        h = d.styles[name]; h.font.name = 'Calibri'; h.font.size = Pt(size); h.font.color.rgb = ACC; h.font.bold = True
        h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    # footer page number
    p = s.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.text = 'PAGE'; f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    r._r.append(f1); r._r.append(it); r._r.append(f2); r.font.size = Pt(9); r.font.color.rgb = MUT
    return d

def para(d, text, bold=False, italic=False, size=None, color=None, align=None, after=None, indent=None):
    p = d.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if indent is not None: p.paragraph_format.left_indent = Cm(indent)
    return p

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor); tcPr.append(sh)

def keep_with_next(p): p.paragraph_format.keep_with_next = True

# ---------------------------------------------------------------- Word paper
def build_paper():
    d = base_doc()
    para(d, 'AR GAUNTLET SECTIONAL', bold=True, size=24, color=ACC, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(d, 'Logical Reasoning: Analytical Reasoning only', size=13, color=MUT, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    t = d.add_table(rows=0, cols=2); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for a, b in [('Questions', '30 (6 sets × 5)'), ('Time', 'No limit. Note your time.'),
                 ('Marking', '+1 right, −0.25 wrong, 0 blank'), ('Pattern', 'CLAT 2026 Section IV formats, pitched well above CLAT')]:
        c = t.add_row().cells; c[0].text = a; c[1].text = b; shade(c[0], 'E3ECF5')
        c[0].paragraphs[0].runs[0].bold = True
    para(d, '', after=6)
    d.add_heading('The six sets', level=2)
    t = d.add_table(rows=1, cols=3); t.style = 'Table Grid'
    for i, h in enumerate(['Set', 'Type', 'Questions']):
        t.rows[0].cells[i].text = h; shade(t.rows[0].cells[i], 'E3ECF5'); t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    n = 1
    for i, p in enumerate(P):
        c = t.add_row().cells; c[0].text = ROMAN[i]; c[1].text = p['kind']; c[2].text = f'{n}–{n+4}'; n += 5
    d.add_heading('Conventions for every set', level=2)
    for c in CONVENTIONS: d.add_paragraph(c, style='List Bullet')
    d.add_heading('Instructions', level=2)
    for c in ['Every set has exactly one correct option per question. Some sets allow more than one valid arrangement; read each question with that in mind.',
              'A question that adds, removes or changes a condition does so for that question only.',
              'Every set opens with a hard "odd one out" question.',
              'Mark answers on the bubble sheet at the end. No answers appear anywhere in this booklet.']:
        d.add_paragraph(c, style='List Bullet')
    n = 1
    for i, p in enumerate(P):
        d.add_page_break()
        h = d.add_heading(f"{p['title']}  (Questions {n}–{n+4})", level=1)
        para(d, p['kind'].upper(), size=8.5, color=MUT, after=6)
        para(d, p['intro'], after=6)
        for j, c in enumerate(p['conds'], 1):
            q = para(d, f'{j}.  {c}', after=2, indent=0.4)
            q.paragraph_format.first_line_indent = Cm(-0.4)
        para(d, '', after=4)
        for q in p['qs']:
            lines = q['stem'].split('\n')
            s = d.add_paragraph(); keep_with_next(s)
            r = s.add_run(f'{n}. '); r.bold = True; r.font.color.rgb = ACC
            s.add_run(lines[0]); s.paragraph_format.space_before = Pt(8)
            for ln in lines[1:]:
                x = para(d, ln, after=1, indent=0.6); keep_with_next(x)
            for k, o in enumerate(q['opts']):
                x = para(d, f'({LET[k]})  {o}', after=1, indent=0.6)
                if k < 3: keep_with_next(x)
            n += 1
    # answer sheet
    d.add_page_break()
    d.add_heading('Answer sheet', level=1)
    para(d, 'Fill one bubble per question. Leave it blank to skip (no penalty).', color=MUT, after=8)
    t = d.add_table(rows=1, cols=6); t.style = 'Table Grid'
    for c in range(6):
        t.rows[0].cells[c].text = f'Set {ROMAN[c]}'; shade(t.rows[0].cells[c], 'E3ECF5'); t.rows[0].cells[c].paragraphs[0].runs[0].bold = True
    for r in range(5):
        cells = t.add_row().cells
        for c in range(6):
            cells[c].text = f'{c*5+r+1:>2}   ○A ○B ○C ○D'
            cells[c].paragraphs[0].runs[0].font.size = Pt(9.5)
    para(d, '', after=10)
    for lab in ['Right: ______', 'Wrong: ______', 'Skipped: ______', 'Score (+1/−0.25): ______', 'Time taken: ______']:
        para(d, lab, after=6)
    d.save(os.path.join(OUT, 'AR_Gauntlet_Sectional_Paper.docx'))

# ---------------------------------------------------------------- Word key
def build_key(audit=None):
    d = base_doc()
    para(d, 'AR GAUNTLET SECTIONAL: ANSWER KEY AND SOLUTIONS', bold=True, size=18, color=ACC, after=4)
    para(d, 'Every answer here was checked by brute force: a program listed every arrangement allowed by each set\'s conditions and tested every option of every question against all of them.', color=MUT, after=10)
    d.add_heading('Answers at a glance', level=2)
    t = d.add_table(rows=1, cols=6); t.style = 'Table Grid'
    for c in range(6):
        t.rows[0].cells[c].text = f'Set {ROMAN[c]}'; shade(t.rows[0].cells[c], 'E3ECF5'); t.rows[0].cells[c].paragraphs[0].runs[0].bold = True
    for r in range(5):
        cells = t.add_row().cells
        for c in range(6):
            q = P[c]['qs'][r]; cells[c].text = f'{c*5+r+1}.  {LET[q["ans"]]}'
    if audit:
        d.add_heading('What CLAT 2026 AR looked like, and how this paper goes further', level=2)
        for line in audit['analysis']: d.add_paragraph(line, style='List Bullet')
        d.add_heading('How the paper was tested', level=2)
        for line in audit['audit']: d.add_paragraph(line, style='List Bullet')
    n = 1
    for i, p in enumerate(P):
        d.add_page_break()
        d.add_heading(f"{p['title']}  (Questions {n}–{n+4})", level=1)
        para(d, p['kind'].upper(), size=8.5, color=MUT, after=6)
        d.add_heading('Solution', level=3)
        para(d, p['solution'], after=6)
        d.add_heading('Key steps', level=3)
        for s in p['key_steps']: d.add_paragraph(s, style='List Number')
        for q in p['qs']:
            d.add_heading(f'Q{n}.  Answer ({LET[q["ans"]]}) {q["opts"][q["ans"]]}', level=3)
            para(d, q['stem'].replace('\n', ' '), italic=True, color=MUT, after=3)
            x = d.add_paragraph(); r = x.add_run('Why. '); r.bold = True; x.add_run(q['why'])
            x = d.add_paragraph(); r = x.add_run('The trap. '); r.bold = True; x.add_run(q['trap'])
            n += 1
    d.save(os.path.join(OUT, 'AR_Gauntlet_Sectional_Answer_Key.docx'))

if __name__ == '__main__':
    audit = json.load(open('audit.json')) if os.path.exists('audit.json') else None
    build_html(); build_paper(); build_key(audit)
    print('built', os.listdir(OUT))
